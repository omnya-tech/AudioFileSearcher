import wx
import pygame
import os
import json
import subprocess
import platform
from core.i18n import LocalizationManager
from core.settings import SettingsManager
from core.audio_processor import TranscriptionThread, EVT_RESULT, WEAK_WORD_THRESHOLD
from core.cross_file_search import AUDIO_EXTENSIONS, parse_subtitles, find_audio_for
from core.time_utils import format_range, format_srt_time
from core.logger import log_error
from gui.settings_dialog import SettingsDialog
from gui.download_dialog import DownloadDialog
from gui.about_dialog import AboutDialog
from gui.report_dialog import ReportDialog
from gui.history_dialog import HistoryDialog
from gui.custom_dict_dialog import CustomDictDialog
from gui.audio_player import AudioPlayerPanel
from gui.search_panel import CrossFileSearchPanel
from gui.processing_dialog import ProcessingDialog
from gui.edit_segment_dialog import EditSegmentDialog

EXPORT_FORMATS = ["srt", "txt", "vtt", "json", "docx"]
EXPORT_WILDCARDS = {
    "srt": "SRT (*.srt)|*.srt",
    "txt": "Text (*.txt)|*.txt",
    "vtt": "WebVTT (*.vtt)|*.vtt",
    "json": "JSON (*.json)|*.json",
    "docx": "Word (*.docx)|*.docx",
}


def is_audio_file(path):
    return path.lower().endswith(AUDIO_EXTENSIONS)


class AudioDropTarget(wx.FileDropTarget):
    def __init__(self, window):
        super().__init__()
        self.window = window

    def OnDropFiles(self, x, y, filenames):
        wx.CallAfter(self.window.on_files_dropped, filenames)
        return True

class MainWindow(wx.Frame):
    def __init__(self, i18n: LocalizationManager, settings: SettingsManager, parent=None):
        title = f"{i18n.get('app_name')} - {i18n.get('app_version')}"
        super().__init__(parent, title=title, size=(950, 700))
        self.i18n = i18n
        i18n.apply_direction(self)
        self.settings = settings
        self.audio_path = None
        self.all_segments = []
        self.displayed_indices = []
        self.unsaved_edits = False
        self.current_percent = None
        self.is_batch_mode = False
        self.batch_queue = []
        self.batch_current_idx = 0
        self.batch_failed = 0
        self.download_dialog = None
        self.processing_dialog = None
        self.transcription_thread = None

        try:
            pygame.mixer.init()
        except Exception as e:
            # لا توجد كارت صوت أو جهاز تشغيل: البرنامج يعمل، والتشغيل فقط هو المعطل
            log_error(f"Audio device init failed: {e}")
        self.i18n.add_observer(self.refresh_ui_texts)

        self.setup_menu()
        self.setup_ui()
        self.setup_accessibility()

        self.apply_theme(self.settings.get("theme", "light"))
        self.apply_font_size(self.settings.get("font_size", 10))

        EVT_RESULT(self, self.on_transcription_update)

        self.Center()
        wx.CallAfter(self.update_title_with_tab)

    # ------------------------------------------------------------------ العنوان والتبويبات
    def get_base_title(self):
        return f"{self.i18n.get('app_name')} - {self.i18n.get('app_version')}"

    def update_title_with_tab(self):
        if hasattr(self, 'notebook'):
            sel = self.notebook.GetSelection()
            if sel != wx.NOT_FOUND:
                tab_name = self.notebook.GetPageText(sel)
                self.SetTitle(f"{self.get_base_title()} - {tab_name}")
                status_template = self.i18n.get("status_current_tab")
                if "{0}" in status_template:
                    self.status_bar.SetStatusText(status_template.format(tab_name))
                else:
                    self.status_bar.SetStatusText(f"[{tab_name}]")

    def on_tab_changed(self, event):
        self.update_title_with_tab()
        event.Skip()

    def on_key_press(self, event):
        keycode = event.GetKeyCode()
        ctrl_down = event.ControlDown()
        shift_down = event.ShiftDown()
        if ctrl_down and keycode == wx.WXK_TAB:
            self.switch_tab(-1 if shift_down else 1)
            return
        if ctrl_down and keycode == wx.WXK_RIGHT:
            self.switch_tab(1)
            return
        elif ctrl_down and keycode == wx.WXK_LEFT:
            self.switch_tab(-1)
            return
        event.Skip()

    def switch_tab(self, direction):
        if hasattr(self, 'notebook'):
            count = self.notebook.GetPageCount()
            if count > 1:
                current = self.notebook.GetSelection()
                next_tab = (current + direction) % count
                self.notebook.SetSelection(next_tab)
                self.update_title_with_tab()

    # ------------------------------------------------------------------ القوائم
    def _menu_labels(self):
        """نصوص عناصر القوائم مع اختصاراتها (تُستخدم عند الإنشاء وعند تغيير اللغة)"""
        return {
            "mi_select_audio": self.i18n.get("btn_select_audio") + "\tCtrl+O",
            "mi_select_folder": self.i18n.get("btn_select_folder") + "\tCtrl+M",
            "mi_open_srt": self.i18n.get("menu_open_srt") + "\tCtrl+Shift+O",
            "mi_export": self.i18n.get("menu_export_srt") + "\tCtrl+S",
            "mi_exit": self.i18n.get("menu_exit") + "\tAlt+F4",
            "mi_stop_audio": self.i18n.get("menu_stop_audio") + "\tF4",
            "mi_check_progress": self.i18n.get("menu_check_progress") + "\tCtrl+I",
            "mi_history": self.i18n.get("menu_history") + "\tCtrl+H",
            "mi_settings": self.i18n.get("menu_settings") + "\tCtrl+P",
            "mi_download_model": self.i18n.get("menu_download_model") + "\tCtrl+D",
            "mi_custom_dict": self.i18n.get("btn_custom_dictionary") + "\tCtrl+K",
            "mi_about": self.i18n.get("menu_about") + "\tF1",
        }

    def setup_menu(self):
        labels = self._menu_labels()
        menubar = wx.MenuBar()
        file_menu = wx.Menu()
        self.mi_select_audio = file_menu.Append(wx.ID_ANY, labels["mi_select_audio"])
        self.mi_select_folder = file_menu.Append(wx.ID_ANY, labels["mi_select_folder"])
        self.mi_open_srt = file_menu.Append(wx.ID_ANY, labels["mi_open_srt"])
        file_menu.AppendSeparator()
        self.mi_export = file_menu.Append(wx.ID_ANY, labels["mi_export"])
        file_menu.AppendSeparator()
        self.mi_exit = file_menu.Append(wx.ID_EXIT, labels["mi_exit"])

        view_menu = wx.Menu()
        self.mi_stop_audio = view_menu.Append(wx.ID_ANY, labels["mi_stop_audio"])
        self.mi_check_progress = view_menu.Append(wx.ID_ANY, labels["mi_check_progress"])
        view_menu.AppendSeparator()
        self.mi_history = view_menu.Append(wx.ID_ANY, labels["mi_history"])

        tools_menu = wx.Menu()
        self.mi_settings = tools_menu.Append(wx.ID_ANY, labels["mi_settings"])
        tools_menu.AppendSeparator()
        self.mi_download_model = tools_menu.Append(wx.ID_ANY, labels["mi_download_model"])
        self.mi_custom_dict = tools_menu.Append(wx.ID_ANY, labels["mi_custom_dict"])

        help_menu = wx.Menu()
        self.mi_about = help_menu.Append(wx.ID_ANY, labels["mi_about"])

        menubar.Append(file_menu, self.i18n.get("menu_file"))
        menubar.Append(view_menu, self.i18n.get("menu_view"))
        menubar.Append(tools_menu, self.i18n.get("menu_tools"))
        menubar.Append(help_menu, self.i18n.get("menu_help"))
        self.SetMenuBar(menubar)

        self.Bind(wx.EVT_MENU, self.on_select_audio, self.mi_select_audio)
        self.Bind(wx.EVT_MENU, self.on_select_folder, self.mi_select_folder)
        self.Bind(wx.EVT_MENU, self.on_open_srt, self.mi_open_srt)
        self.Bind(wx.EVT_MENU, self.on_export_srt, self.mi_export)
        self.Bind(wx.EVT_MENU, self.open_settings, self.mi_settings)
        self.Bind(wx.EVT_MENU, self.on_stop_audio, self.mi_stop_audio)
        self.Bind(wx.EVT_MENU, self.on_check_progress, self.mi_check_progress)
        self.Bind(wx.EVT_MENU, self.on_open_download_dialog, self.mi_download_model)
        self.Bind(wx.EVT_MENU, self.on_open_custom_dict, self.mi_custom_dict)
        self.Bind(wx.EVT_MENU, self.on_open_history, self.mi_history)
        self.Bind(wx.EVT_MENU, self.on_about, self.mi_about)
        self.Bind(wx.EVT_MENU, self.on_exit, self.mi_exit)

        self.mi_export.Enable(False)

    # ------------------------------------------------------------------ الواجهة
    def setup_ui(self):
        self.panel = wx.Panel(self)
        self.panel.SetDropTarget(AudioDropTarget(self))
        main_sizer = wx.BoxSizer(wx.VERTICAL)

        self.notebook = wx.Notebook(self.panel)
        self.notebook.Bind(wx.EVT_NOTEBOOK_PAGE_CHANGED, self.on_tab_changed)

        self.transcription_panel = wx.Panel(self.notebook)
        self.transcription_panel.SetDropTarget(AudioDropTarget(self))
        trans_sizer = wx.BoxSizer(wx.VERTICAL)

        file_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_select = wx.Button(self.transcription_panel, label=self.i18n.get("btn_select_audio"))
        self.btn_select_folder = wx.Button(self.transcription_panel, label=self.i18n.get("btn_select_folder"))
        self.txt_file_path = wx.TextCtrl(self.transcription_panel, style=wx.TE_READONLY)

        file_sizer.Add(self.btn_select, 0, wx.ALL | wx.ALIGN_CENTER_VERTICAL, 5)
        file_sizer.Add(self.btn_select_folder, 0, wx.ALL | wx.ALIGN_CENTER_VERTICAL, 5)
        file_sizer.Add(self.txt_file_path, 1, wx.ALL | wx.EXPAND, 5)
        trans_sizer.Add(file_sizer, 0, wx.EXPAND | wx.ALL, 5)

        btn_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_process = wx.Button(self.transcription_panel, label=self.i18n.get("btn_process"))
        self.btn_export = wx.Button(self.transcription_panel, label=self.i18n.get("btn_export"))
        self.btn_process.Disable()
        self.btn_export.Disable()
        btn_sizer.Add(self.btn_process, 0, wx.ALL, 5)
        btn_sizer.Add(self.btn_export, 0, wx.ALL, 5)
        trans_sizer.Add(btn_sizer, 0, wx.CENTER | wx.ALL, 5)

        filter_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.lbl_filter = wx.StaticText(self.transcription_panel, label=self.i18n.get("lbl_search"))
        self.txt_filter = wx.TextCtrl(self.transcription_panel, style=wx.TE_PROCESS_ENTER)
        filter_sizer.Add(self.lbl_filter, 0, wx.ALL | wx.ALIGN_CENTER_VERTICAL, 5)
        filter_sizer.Add(self.txt_filter, 1, wx.ALL | wx.EXPAND, 5)
        trans_sizer.Add(filter_sizer, 0, wx.EXPAND | wx.ALL, 5)

        self.result_list = wx.ListCtrl(self.transcription_panel, style=wx.LC_REPORT | wx.LC_SINGLE_SEL)
        self._insert_result_columns()
        trans_sizer.Add(self.result_list, 1, wx.EXPAND | wx.ALL, 10)

        self.transcription_panel.SetSizer(trans_sizer)

        self.search_panel = CrossFileSearchPanel(self.notebook, self.i18n, play_callback=self.play_audio_from_search)

        self.notebook.AddPage(self.transcription_panel, self.i18n.get("tab_transcription"))
        self.notebook.AddPage(self.search_panel, self.i18n.get("tab_global_search"))
        self.i18n.fix_notebook(self.notebook)

        main_sizer.Add(self.notebook, 1, wx.EXPAND | wx.ALL, 5)

        self.audio_player = AudioPlayerPanel(self.panel, self.i18n)
        self.audio_player.Hide()
        main_sizer.Add(self.audio_player, 0, wx.EXPAND | wx.ALL, 5)

        self.status_bar = self.CreateStatusBar()
        self.status_bar.SetStatusText(self.i18n.get("status_ready"))

        self.panel.SetSizer(main_sizer)

        self.btn_select.Bind(wx.EVT_BUTTON, self.on_select_audio)
        self.btn_select_folder.Bind(wx.EVT_BUTTON, self.on_select_folder)
        self.btn_process.Bind(wx.EVT_BUTTON, self.on_process)
        self.btn_export.Bind(wx.EVT_BUTTON, self.on_export_srt)
        self.txt_filter.Bind(wx.EVT_TEXT, self.on_filter_results)
        self.result_list.Bind(wx.EVT_LIST_ITEM_ACTIVATED, self.on_play_segment)
        self.result_list.Bind(wx.EVT_KEY_DOWN, self.on_list_key)
        self.result_list.Bind(wx.EVT_CONTEXT_MENU, self.on_list_context_menu)
        self.Bind(wx.EVT_CLOSE, self.on_exit)

        self.setup_button_hover_effects()

    def _insert_result_columns(self):
        self.result_list.InsertColumn(0, self.i18n.get("list_header_time"), width=150)
        self.result_list.InsertColumn(1, self.i18n.get("list_header_text"), width=700)

    def setup_accessibility(self):
        self.Bind(wx.EVT_CHAR_HOOK, self.on_key_press)
        # اختصار Ctrl+F لا يوجد في القوائم، فنربطه بمعرف خاص ينقل التركيز لحقل البحث
        self.id_focus_filter = wx.NewIdRef()
        self.Bind(wx.EVT_MENU, self.on_focus_filter, id=self.id_focus_filter)
        accel_tbl = wx.AcceleratorTable([
            (wx.ACCEL_CTRL, ord('F'), self.id_focus_filter),
        ])
        self.SetAcceleratorTable(accel_tbl)
        self.btn_select.SetFocus()

    def on_focus_filter(self, event):
        self.notebook.SetSelection(0)
        self.txt_filter.SetFocus()
        self.txt_filter.SelectAll()

    def setup_button_hover_effects(self):
        buttons = [self.btn_select, self.btn_select_folder, self.btn_process, self.btn_export]
        for btn in buttons:
            btn.Bind(wx.EVT_ENTER_WINDOW, self.on_button_hover)
            btn.Bind(wx.EVT_LEAVE_WINDOW, self.on_button_leave)

    def on_button_hover(self, event):
        btn = event.GetEventObject()
        if btn.IsEnabled():
            theme = self.settings.get("theme", "light")
            if theme == "dark": btn.SetBackgroundColour(wx.Colour(45, 56, 78))
            else: btn.SetBackgroundColour(wx.Colour(233, 236, 239))
            btn.Refresh()
        event.Skip()

    def on_button_leave(self, event):
        btn = event.GetEventObject()
        theme = self.settings.get("theme", "light")
        if theme == "dark": btn.SetBackgroundColour(wx.Colour(30, 38, 54))
        else: btn.SetBackgroundColour(wx.Colour(255, 255, 255))
        btn.Refresh()
        event.Skip()

    # ------------------------------------------------------------------ النوافذ الفرعية
    def open_settings(self, event):
        dlg = SettingsDialog(self, self.i18n, self.settings)
        dlg.ShowModal()
        dlg.Destroy()

    def on_open_download_dialog(self, event):
        if not getattr(self, 'download_dialog', None):
            self.download_dialog = DownloadDialog(self, self.i18n)
            self.download_dialog.Show()
        else:
            if self.download_dialog.IsIconized(): self.download_dialog.Restore()
            self.download_dialog.Show()
            self.download_dialog.Raise()

    def on_open_custom_dict(self, event):
        dlg = CustomDictDialog(self, self.i18n, self.settings)
        dlg.ShowModal()
        dlg.Destroy()

    def on_about(self, event):
        dlg = AboutDialog(self, self.i18n)
        dlg.ShowModal()
        dlg.Destroy()

    def on_open_history(self, event):
        dlg = HistoryDialog(self, self.i18n)
        dlg.ShowModal()
        dlg.Destroy()

    def on_check_progress(self, event):
        if self.current_percent is not None:
            msg_template = self.i18n.get("msg_progress_info")
            if "{percent}" in msg_template:
                msg = msg_template.format(percent=self.current_percent)
            else:
                msg = f"Progress: {self.current_percent}%"
            wx.MessageBox(msg, self.i18n.get("dialog_progress_title"), wx.ICON_INFORMATION)
        else:
            wx.MessageBox(self.i18n.get("msg_no_active_process"), self.i18n.get("dialog_progress_title"), wx.ICON_INFORMATION)

    def show_error(self, error):
        msg_temp = self.i18n.get("status_error")
        msg = msg_temp.format(error=str(error)) if "{error}" in msg_temp else str(error)
        wx.MessageBox(msg, self.i18n.get("dialog_error_title"), wx.ICON_ERROR)

    # ------------------------------------------------------------------ فتح الملفات
    def on_files_dropped(self, filenames):
        if not filenames or self.transcription_thread_running(): return
        path = filenames[0]
        if os.path.isdir(path):
            self.start_folder(path, auto_start=True)
        elif os.path.isfile(path):
            if is_audio_file(path):
                self.set_single_file(path)
                self.on_process(None)
            elif path.lower().endswith(('.srt', '.vtt')):
                self.load_srt_file(path)

    def transcription_thread_running(self):
        return bool(self.transcription_thread and self.transcription_thread.is_alive() and not self.transcription_thread.aborted)

    def set_single_file(self, path):
        self.is_batch_mode = False
        self.batch_queue = []
        self.audio_path = path
        self.txt_file_path.SetValue(path)
        self.btn_process.Enable()

    def start_folder(self, folder_path, auto_start=False):
        files = sorted(os.path.join(folder_path, f) for f in os.listdir(folder_path)
                       if is_audio_file(f) and os.path.isfile(os.path.join(folder_path, f)))
        if not files:
            wx.MessageBox(self.i18n.get("msg_folder_no_audio"), self.i18n.get("dialog_error_title"), wx.ICON_WARNING)
            return
        self.batch_queue = files
        self.is_batch_mode = True
        self.audio_path = folder_path
        self.txt_file_path.SetValue(folder_path)
        self.btn_process.Enable()
        if auto_start:
            self.on_process(None)
        else:
            self.txt_file_path.SetFocus()

    def play_audio_from_search(self, audio_path, start_time):
        self._play(audio_path, start_time)

    def on_open_srt(self, event):
        wildcard = "Subtitles (*.srt;*.vtt)|*.srt;*.vtt|All Files (*.*)|*.*"
        dlg = wx.FileDialog(self, message=self.i18n.get("dialog_open_srt"), wildcard=wildcard, style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST)
        if dlg.ShowModal() == wx.ID_OK:
            self.load_srt_file(dlg.GetPath())
        dlg.Destroy()

    def load_srt_file(self, srt_path):
        if not self.confirm_discard_edits():
            return
        try:
            parsed = parse_subtitles(srt_path)
        except Exception as e:
            log_error(f"Failed to open subtitles {srt_path}: {e}")
            self.show_error(e)
            return
        if not parsed:
            wx.MessageBox(self.i18n.get("msg_srt_empty"), self.i18n.get("dialog_warning_title"), wx.ICON_WARNING)
            return

        self.all_segments = [(format_range(s, e), text, s, e, []) for s, e, text in parsed]
        self.txt_filter.ChangeValue("")
        self.update_list()
        self.txt_file_path.SetValue(srt_path)
        self.btn_export.Enable()
        self.mi_export.Enable(True)

        self.is_batch_mode = False
        self.batch_queue = []
        self.audio_path = find_audio_for(srt_path)
        self.btn_process.Enable(bool(self.audio_path))
        self.status_bar.SetStatusText(self.i18n.get("status_ready"))

    def on_select_audio(self, event):
        exts = ";".join(f"*{e}" for e in AUDIO_EXTENSIONS)
        wildcard = f"Audio ({exts})|{exts}|All Files (*.*)|*.*"
        dlg = wx.FileDialog(self, message=self.i18n.get("dialog_open_audio"), wildcard=wildcard, style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST)
        if dlg.ShowModal() == wx.ID_OK:
            self.set_single_file(dlg.GetPath())
            self.txt_file_path.SetFocus()
        dlg.Destroy()

    def on_select_folder(self, event):
        dlg = wx.DirDialog(self, message=self.i18n.get("dialog_select_folder"), style=wx.DD_DEFAULT_STYLE | wx.DD_DIR_MUST_EXIST)
        if dlg.ShowModal() == wx.ID_OK:
            self.start_folder(dlg.GetPath())
        dlg.Destroy()

    # ------------------------------------------------------------------ التفريغ
    def _set_controls_busy(self, busy):
        self.btn_select.Enable(not busy)
        self.btn_select_folder.Enable(not busy)
        self.btn_process.Enable(not busy and bool(self.audio_path or self.batch_queue))
        has_results = bool(self.all_segments)
        self.btn_export.Enable(not busy and has_results)
        self.mi_export.Enable(not busy and has_results)
        self.mi_select_audio.Enable(not busy)
        self.mi_select_folder.Enable(not busy)
        self.mi_open_srt.Enable(not busy)

    def _close_processing_dialog(self):
        if self.processing_dialog:
            self.processing_dialog.Destroy()
            self.processing_dialog = None

    def on_process(self, event):
        if not self.audio_path and not self.batch_queue: return
        if self.transcription_thread_running(): return
        if not self.confirm_discard_edits(): return
        if not self.is_batch_mode and not os.path.isfile(self.audio_path or ""):
            wx.MessageBox(self.i18n.get("msg_audio_not_found"), self.i18n.get("dialog_error_title"), wx.ICON_ERROR)
            return

        self._set_controls_busy(True)
        self.result_list.SetFocus()

        self._close_processing_dialog()
        self.processing_dialog = ProcessingDialog(self, self.i18n)
        self.processing_dialog.Show()

        if self.is_batch_mode:
            self.batch_current_idx = 0
            self.batch_failed = 0
            self.process_next_in_batch()
        else:
            self.all_segments = []
            self.result_list.DeleteAllItems()
            self.current_percent = 0
            self.on_stop_audio(None)

            title_template = self.i18n.get("window_title_progress")
            if "{percent}" in title_template and "{app_title}" in title_template:
                self.SetTitle(title_template.format(percent=0, app_title=self.get_base_title()))
            else:
                self.SetTitle(self.get_base_title())

            filename = os.path.basename(self.audio_path)
            self.processing_dialog.update_progress(0, 100, self.i18n.get("status_init_engine"), filename)
            self.transcription_thread = TranscriptionThread(self, self.audio_path, self.i18n)

    def process_next_in_batch(self):
        if self.batch_current_idx < len(self.batch_queue):
            current_file = self.batch_queue[self.batch_current_idx]
            self.audio_path = current_file
            self.txt_file_path.SetValue(current_file)
            self.all_segments = []
            self.result_list.DeleteAllItems()
            self.current_percent = 0
            self.on_stop_audio(None)

            filename = os.path.basename(current_file)
            self.status_bar.SetStatusText(f"[{self.batch_current_idx + 1}/{len(self.batch_queue)}] {filename}")

            title_template = self.i18n.get("window_title_batch_progress")
            if "{current}" in title_template and "{total}" in title_template:
                self.SetTitle(title_template.format(current=self.batch_current_idx + 1, total=len(self.batch_queue), percent=0, app_title=self.get_base_title()))
            else:
                self.SetTitle(self.get_base_title())

            if self.processing_dialog:
                self.processing_dialog.update_progress(0, 100, self.i18n.get("status_init_engine"), filename)

            self.transcription_thread = TranscriptionThread(self, current_file, self.i18n)
        else:
            total = len(self.batch_queue)
            succeeded = total - self.batch_failed
            out_dir = self._output_dir_for(self.batch_queue[0]) if self.batch_queue else None
            self.is_batch_mode = False
            self.batch_queue = []
            self.update_title_with_tab()
            self.status_bar.SetStatusText(self.i18n.get("status_ready"))
            self._set_controls_busy(False)
            self._close_processing_dialog()

            msg = self.i18n.get("msg_batch_done", count=succeeded)
            if self.batch_failed:
                msg += "\n" + self.i18n.get("msg_batch_failed_count", count=self.batch_failed)
            wx.MessageBox(msg, self.i18n.get("dialog_success_title"), wx.ICON_INFORMATION)
            if out_dir and succeeded and self.settings.get("open_folder_after_save", False):
                self.open_in_file_manager(out_dir)

    def _output_dir_for(self, audio_path):
        output_dir = self.settings.get("output_directory", "")
        return output_dir if output_dir and os.path.isdir(output_dir) else os.path.dirname(audio_path)

    def auto_save_results(self):
        """حفظ النتائج تلقائياً. في وضع المجلد يتم الحفظ دائماً حتى لا تضيع نتائج الملفات السابقة"""
        if not (self.settings.get("auto_save", True) or self.is_batch_mode):
            return None
        default_format = self.settings.get("default_export_format", "srt")
        if default_format not in EXPORT_FORMATS: default_format = "srt"
        base_name = os.path.splitext(os.path.basename(self.audio_path))[0]
        save_path = os.path.join(self._output_dir_for(self.audio_path), f"{base_name}.{default_format}")
        return save_path if self.save_as(save_path, default_format, show_msg=False) else None

    def open_in_file_manager(self, path):
        try:
            if platform.system() == "Windows":
                if os.path.isfile(path): subprocess.Popen(['explorer', '/select,', os.path.normpath(path)])
                else: os.startfile(os.path.normpath(path))
            elif platform.system() == "Darwin":
                subprocess.Popen(['open', '-R', path] if os.path.isfile(path) else ['open', path])
            else:
                subprocess.Popen(['xdg-open', os.path.dirname(path) if os.path.isfile(path) else path])
        except Exception as e:
            log_error(f"Failed to open file manager: {e}")

    def on_transcription_update(self, event):
        # تجاهل أي حدث قادم من عملية قديمة تم إلغاؤها
        if event.source is not None and event.source is not self.transcription_thread:
            return
        if self.transcription_thread is not None and self.transcription_thread.aborted:
            return

        status = event.status
        filename = os.path.basename(self.audio_path) if self.audio_path else ""

        if status == "loading" or status == "loading_local":
            status_txt = self.i18n.get("status_loading_model") if status == "loading" else self.i18n.get("status_loading_local_model")
            self.status_bar.SetStatusText(status_txt)
            if self.processing_dialog:
                self.processing_dialog.update_progress(0, 100, status_txt, filename)

        elif status == "transcribing":
            if self.processing_dialog:
                self.processing_dialog.update_progress(0, 100, self.i18n.get("status_extracting"), filename)

        elif status == "progress":
            self.current_percent = event.data
            if self.processing_dialog:
                self.processing_dialog.update_progress(event.data, 100, self.i18n.get("status_extracting"), filename)

            if self.is_batch_mode:
                title_template = self.i18n.get("window_title_batch_progress")
                if "{current}" in title_template:
                    self.SetTitle(title_template.format(current=self.batch_current_idx + 1, total=len(self.batch_queue), percent=event.data, app_title=self.get_base_title()))
            else:
                title_template = self.i18n.get("window_title_progress")
                if "{percent}" in title_template:
                    self.SetTitle(title_template.format(percent=event.data, app_title=self.get_base_title()))

            status_template = self.i18n.get("status_transcribing")
            if "{percent}" in status_template:
                self.status_bar.SetStatusText(status_template.format(percent=event.data))

        elif status == "done":
            self.current_percent = None
            self.transcription_thread = None
            self.all_segments = event.data["results"]
            self.txt_filter.ChangeValue("")
            self.update_list()

            status_template = self.i18n.get("status_done")
            status_txt = status_template.format(time=event.data["time"]) if "{time}" in status_template else "Done"

            if self.processing_dialog:
                self.processing_dialog.update_progress(100, 100, status_txt, filename)

            if not self.all_segments:
                # لا يوجد كلام في الملف: لا داعي لحفظ ملف فارغ
                saved_path = None
            else:
                saved_path = self.auto_save_results()

            if self.is_batch_mode:
                self.batch_current_idx += 1
                self.process_next_in_batch()
                return

            self.update_title_with_tab()
            self.status_bar.SetStatusText(status_txt)
            self._set_controls_busy(False)
            self.txt_filter.SetFocus()
            self._close_processing_dialog()

            if saved_path and self.settings.get("open_folder_after_save", False):
                self.open_in_file_manager(saved_path)

            if not self.all_segments:
                wx.MessageBox(self.i18n.get("report_no_segments"), self.i18n.get("dialog_info_title"), wx.ICON_INFORMATION)
            elif event.data.get("report"):
                report_dlg = ReportDialog(self, self.i18n, event.data["report"])
                report_dlg.ShowModal()
                report_dlg.Destroy()

        elif status == "error":
            self.current_percent = None
            self.transcription_thread = None

            if self.is_batch_mode:
                log_error(f"Batch item failed {filename}: {event.data}")
                self.batch_failed += 1
                msg_template = self.i18n.get("status_batch_error")
                if "{file}" in msg_template:
                    self.status_bar.SetStatusText(msg_template.format(file=filename, error=event.data))
                self.batch_current_idx += 1
                self.process_next_in_batch()
            else:
                self.update_title_with_tab()
                msg_template = self.i18n.get("status_error")
                msg = msg_template.format(error=event.data) if "{error}" in msg_template else f"Error: {event.data}"
                self.status_bar.SetStatusText(msg)
                self._close_processing_dialog()
                self._set_controls_busy(False)
                wx.MessageBox(msg, self.i18n.get("dialog_error_title"), wx.ICON_ERROR)
                self.result_list.SetFocus()

    def cancel_processing(self):
        if self.transcription_thread:
            self.transcription_thread.abort()
        self.transcription_thread = None
        if self.is_batch_mode:
            self.audio_path = None
            self.txt_file_path.SetValue("")
        self.is_batch_mode = False
        self.batch_queue = []
        self.current_percent = None
        self.update_title_with_tab()
        self._close_processing_dialog()
        self._set_controls_busy(False)
        self.status_bar.SetStatusText(self.i18n.get("status_canceled"))

    # ------------------------------------------------------------------ عرض النتائج
    @staticmethod
    def _is_weak(seg):
        """المقطع فيه كلمة واحدة على الأقل النموذج غير متأكد منها"""
        words = seg[4] if len(seg) > 4 and seg[4] else []
        return any(w.get("probability", 1.0) < WEAK_WORD_THRESHOLD for w in words)

    def _fill_row(self, row, seg_index):
        seg = self.all_segments[seg_index]
        self.result_list.SetItem(row, 1, seg[1].replace("\n", " "))
        colour = wx.Colour(200, 40, 40) if self._is_weak(seg) else self.result_list.GetForegroundColour()
        self.result_list.SetItemTextColour(row, colour)

    def update_list(self, indices=None):
        """عرض المقاطع. indices أرقام المقاطع داخل all_segments (الكل لو None)"""
        if indices is None:
            indices = list(range(len(self.all_segments)))
        self.result_list.Freeze()
        try:
            self.result_list.DeleteAllItems()
            self.displayed_indices = indices
            for seg_index in indices:
                seg = self.all_segments[seg_index]
                row = self.result_list.InsertItem(self.result_list.GetItemCount(), format_range(seg[2], seg[3]))
                # كل سطر يحمل رقم المقطع الأصلي، فالتشغيل والتعديل يعملان حتى أثناء البحث
                self.result_list.SetItemData(row, seg_index)
                self._fill_row(row, seg_index)
        finally:
            self.result_list.Thaw()

    def on_filter_results(self, event):
        query = self.txt_filter.GetValue().strip().lower()
        if not query:
            self.update_list()
            return
        self.update_list([i for i, seg in enumerate(self.all_segments) if query in seg[1].lower()])

    def _selected_segment_index(self):
        row = self.result_list.GetFirstSelected()
        if row == -1:
            return None, None
        seg_index = self.result_list.GetItemData(row)
        return (row, seg_index) if 0 <= seg_index < len(self.all_segments) else (None, None)

    def _can_play(self):
        return not self.is_batch_mode and self.audio_path and os.path.isfile(self.audio_path)

    # ------------------------------------------------------------------ تعديل النص
    def on_list_key(self, event):
        if event.GetKeyCode() == wx.WXK_F2:
            self.on_edit_segment(None)
        else:
            event.Skip()

    def on_list_context_menu(self, event):
        row, seg_index = self._selected_segment_index()
        if seg_index is None:
            return
        menu = wx.Menu()
        mi_play = menu.Append(wx.ID_ANY, self.i18n.get("menu_play_segment"))
        mi_edit = menu.Append(wx.ID_ANY, self.i18n.get("menu_edit_segment") + "\tF2")
        mi_play.Enable(bool(self._can_play()))
        self.Bind(wx.EVT_MENU, lambda e: self._play(self.audio_path, self.all_segments[seg_index][2]), mi_play)
        self.Bind(wx.EVT_MENU, self.on_edit_segment, mi_edit)
        self.result_list.PopupMenu(menu)
        menu.Destroy()

    def on_edit_segment(self, event):
        if self.transcription_thread_running():
            return
        row, seg_index = self._selected_segment_index()
        if seg_index is None:
            return
        seg = self.all_segments[seg_index]
        play = (lambda: self._play(self.audio_path, seg[2])) if self._can_play() else None
        dlg = EditSegmentDialog(self, self.i18n, format_range(seg[2], seg[3]), seg[1], play)
        if dlg.ShowModal() == wx.ID_OK:
            new_text = dlg.get_text()
            if new_text and new_text != seg[1]:
                # توقيتات الكلمات القديمة لم تعد تطابق النص، والمراجعة البشرية تلغي علامة الشك
                self.all_segments[seg_index] = (seg[0], new_text, seg[2], seg[3], [])
                self._fill_row(row, seg_index)
                self.unsaved_edits = True
                self.status_bar.SetStatusText(self.i18n.get("status_segment_edited"))
        dlg.Destroy()
        self.result_list.SetFocus()

    def confirm_discard_edits(self):
        """يُستدعى قبل أي عملية تستبدل النتائج الحالية. يرجع False لو المستخدم تراجع"""
        if not getattr(self, 'unsaved_edits', False):
            return True
        dlg = wx.MessageDialog(self, self.i18n.get("msg_unsaved_edits"), self.i18n.get("dialog_warning_title"),
                               wx.YES_NO | wx.CANCEL | wx.YES_DEFAULT | wx.ICON_QUESTION)
        dlg.SetYesNoCancelLabels(self.i18n.get("btn_save"), self.i18n.get("btn_discard"), self.i18n.get("btn_cancel"))
        res = dlg.ShowModal()
        dlg.Destroy()
        if res == wx.ID_YES:
            self.on_export_srt(None)
            return not self.unsaved_edits
        if res == wx.ID_NO:
            self.unsaved_edits = False
            return True
        return False

    # ------------------------------------------------------------------ التشغيل
    def _play(self, audio_path, start_time):
        if not audio_path or not os.path.isfile(audio_path):
            wx.MessageBox(self.i18n.get("msg_audio_not_found"), self.i18n.get("dialog_error_title"), wx.ICON_WARNING)
            return
        if not self.audio_player.IsShown():
            self.audio_player.Show()
            self.panel.Layout()
        ok, error = self.audio_player.load_and_play(audio_path, start_time or 0)
        if ok:
            self.status_bar.SetStatusText(self.i18n.get("status_playing"))
        else:
            wx.MessageBox(self.i18n.get("msg_playback_failed", error=error), self.i18n.get("dialog_error_title"), wx.ICON_WARNING)

    def on_play_segment(self, event):
        if not self._can_play(): return
        seg_index = self.result_list.GetItemData(event.GetIndex())
        if 0 <= seg_index < len(self.all_segments):
            self._play(self.audio_path, self.all_segments[seg_index][2])

    def on_stop_audio(self, event):
        self.audio_player.on_stop(None)
        if self.audio_player.IsShown():
            self.audio_player.Hide()
            self.panel.Layout()
        self.status_bar.SetStatusText(self.i18n.get("status_ready"))

    # ------------------------------------------------------------------ التصدير
    def on_export_srt(self, event):
        if not self.all_segments: return
        enabled = self.settings.get("export_formats", EXPORT_FORMATS) or EXPORT_FORMATS
        formats = [f for f in EXPORT_FORMATS if f in enabled]
        default_format = self.settings.get("default_export_format", "srt")
        # الصيغة الافتراضية تظهر أولاً في مربع الحفظ
        if default_format in formats:
            formats.remove(default_format)
            formats.insert(0, default_format)
        wildcard = "|".join(EXPORT_WILDCARDS[f] for f in formats)

        default_name = ""
        if self.audio_path and os.path.isfile(self.audio_path):
            default_name = os.path.splitext(os.path.basename(self.audio_path))[0] + "." + formats[0]

        dlg = wx.FileDialog(self, message=self.i18n.get("dialog_save_srt"), defaultFile=default_name,
                            wildcard=wildcard, style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
        if dlg.ShowModal() == wx.ID_OK:
            path = dlg.GetPath()
            ext = os.path.splitext(path)[1].lower().lstrip(".")
            if ext not in EXPORT_FORMATS:
                # المستخدم لم يكتب امتداداً: نعتمد الصيغة المختارة في مربع الحفظ
                ext = formats[dlg.GetFilterIndex()] if 0 <= dlg.GetFilterIndex() < len(formats) else "srt"
                path = f"{path}.{ext}"
            self.save_as(path, ext, show_msg=True)
        dlg.Destroy()

    def save_as(self, path, fmt, show_msg=True):
        writers = {"txt": self._write_txt, "vtt": self._write_vtt, "json": self._write_json,
                   "docx": self._write_docx, "srt": self._write_srt}
        try:
            writers.get(fmt, self._write_srt)(path)
        except ImportError:
            wx.MessageBox(self.i18n.get("msg_docx_error"), self.i18n.get("dialog_error_title"), wx.ICON_ERROR)
            return False
        except Exception as e:
            log_error(f"Export to {path} failed: {e}")
            self.show_error(e)
            return False
        self.unsaved_edits = False
        if show_msg:
            wx.MessageBox(self.i18n.get("msg_export_success"), self.i18n.get("dialog_success_title"), wx.ICON_INFORMATION)
        return True

    def _write_txt(self, path):
        with open(path, 'w', encoding='utf-8') as f:
            f.write("\n\n".join(item[1].strip() for item in self.all_segments))

    def _write_srt(self, path):
        with open(path, 'w', encoding='utf-8') as f:
            for i, item in enumerate(self.all_segments, 1):
                f.write(f"{i}\n{format_srt_time(item[2])} --> {format_srt_time(item[3])}\n{item[1]}\n\n")

    def _write_vtt(self, path):
        with open(path, 'w', encoding='utf-8') as f:
            f.write("WEBVTT\n\n")
            for i, item in enumerate(self.all_segments, 1):
                f.write(f"{i}\n{format_srt_time(item[2], '.')} --> {format_srt_time(item[3], '.')}\n{item[1]}\n\n")

    def _write_json(self, path):
        data = []
        for item in self.all_segments:
            seg_data = {"start": item[2], "end": item[3], "text": item[1]}
            if len(item) > 4 and item[4]: seg_data["words"] = item[4]
            data.append(seg_data)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)

    def _write_docx(self, path):
        import docx
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement

        doc = docx.Document()
        is_rtl = self.i18n.language == "ar"
        heading = doc.add_heading(self.i18n.get("app_name"), 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for item in self.all_segments:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(12)
            if is_rtl:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                # اتجاه الفقرة من اليمين لليسار حتى يظهر النص العربي مرتباً في الوورد
                p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
            run_time = p.add_run(f"[{format_range(item[2], item[3])}]\n")
            run_time.bold = True
            run_time.font.color.rgb = RGBColor(100, 100, 100)
            run_text = p.add_run(item[1])
            if is_rtl:
                run_text.font.rtl = True

        doc.save(path)

    # ------------------------------------------------------------------ الإغلاق
    def on_exit(self, event):
        if not self.confirm_discard_edits():
            if isinstance(event, wx.CloseEvent) and event.CanVeto(): event.Veto()
            return
        if getattr(self, 'download_dialog', None) and self.download_dialog.is_downloading:
            dlg = wx.MessageDialog(self, self.i18n.get("dl_msg_confirm_hide"), self.i18n.get("dialog_warning_title"), wx.YES_NO | wx.NO_DEFAULT | wx.ICON_WARNING)
            res = dlg.ShowModal()
            dlg.Destroy()
            if res != wx.ID_YES:
                if isinstance(event, wx.CloseEvent) and event.CanVeto(): event.Veto()
                return
            if getattr(self.download_dialog, 'download_thread', None): self.download_dialog.download_thread.abort()

        if self.transcription_thread:
            self.transcription_thread.abort()
        self.i18n.remove_observer(self.refresh_ui_texts)
        self.audio_player.cleanup()
        try:
            if pygame.mixer.get_init(): pygame.mixer.quit()
        except Exception:
            pass
        self.Destroy()

    # ------------------------------------------------------------------ اللغة والمظهر
    def apply_layout_direction(self):
        """العربية من اليمين لليسار والإنجليزية من اليسار لليمين، للنافذة وكل ما بداخلها"""
        direction = wx.Layout_RightToLeft if self.i18n.language == "ar" else wx.Layout_LeftToRight
        if self.GetLayoutDirection() == direction:
            return
        def apply(widget):
            widget.SetLayoutDirection(direction)
            for child in widget.GetChildren():
                apply(child)
        self.Freeze()
        try:
            apply(self)
            if self.GetStatusBar(): self.GetStatusBar().SetLayoutDirection(direction)
            self.i18n.fix_notebook(self.notebook)
        finally:
            self.Thaw()

    def refresh_ui_texts(self):
        self.apply_layout_direction()
        if hasattr(self, 'notebook'):
            self.notebook.SetPageText(0, self.i18n.get("tab_transcription"))
            self.notebook.SetPageText(1, self.i18n.get("tab_global_search"))

        menubar = self.GetMenuBar()
        if menubar:
            menubar.SetMenuLabel(0, self.i18n.get("menu_file"))
            menubar.SetMenuLabel(1, self.i18n.get("menu_view"))
            menubar.SetMenuLabel(2, self.i18n.get("menu_tools"))
            menubar.SetMenuLabel(3, self.i18n.get("menu_help"))
            for attr, label in self._menu_labels().items():
                getattr(self, attr).SetItemLabel(label)

        self.btn_select.SetLabel(self.i18n.get("btn_select_audio"))
        self.btn_select_folder.SetLabel(self.i18n.get("btn_select_folder"))
        self.btn_process.SetLabel(self.i18n.get("btn_process"))
        self.btn_export.SetLabel(self.i18n.get("btn_export"))
        self.lbl_filter.SetLabel(self.i18n.get("lbl_search"))

        self.result_list.ClearAll()
        self._insert_result_columns()
        self.update_list(self.displayed_indices)

        self.search_panel.refresh_ui_texts()
        self.audio_player.refresh_ui_texts()

        self.update_title_with_tab()
        self.status_bar.SetStatusText(self.i18n.get("status_ready"))
        self.panel.Layout()
        self.Refresh()

    def apply_theme(self, theme):
        if theme == "dark":
            bg_color = wx.Colour(15, 19, 28)
            panel_bg = wx.Colour(22, 29, 43)
            fg_color = wx.Colour(241, 241, 241)
            list_bg = wx.Colour(10, 13, 18)
            button_bg = wx.Colour(30, 38, 54)
        else:
            bg_color = wx.Colour(248, 249, 250)
            panel_bg = wx.Colour(255, 255, 255)
            fg_color = wx.Colour(33, 37, 41)
            list_bg = wx.Colour(255, 255, 255)
            button_bg = wx.Colour(255, 255, 255)

        self.SetBackgroundColour(bg_color)
        self.SetForegroundColour(fg_color)
        self.panel.SetBackgroundColour(panel_bg)
        self.panel.SetForegroundColour(fg_color)

        def apply_recursive(widget):
            if isinstance(widget, wx.ListCtrl):
                widget.SetBackgroundColour(list_bg)
                widget.SetForegroundColour(fg_color)
            elif isinstance(widget, wx.Button):
                widget.SetBackgroundColour(button_bg)
                widget.SetForegroundColour(fg_color)
            elif isinstance(widget, wx.TextCtrl):
                widget.SetBackgroundColour(list_bg if theme == "dark" else wx.Colour(255, 255, 255))
                widget.SetForegroundColour(fg_color)
            elif isinstance(widget, wx.Notebook):
                widget.SetBackgroundColour(bg_color)
                widget.SetForegroundColour(fg_color)
            else:
                widget.SetBackgroundColour(panel_bg)
                widget.SetForegroundColour(fg_color)
            for child in widget.GetChildren(): apply_recursive(child)

        apply_recursive(self.panel)
        self.Refresh()

    def apply_font_size(self, size):
        try:
            size = int(size)
        except (TypeError, ValueError):
            size = 10
        font = wx.Font(size, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_NORMAL)
        def set_font_recursive(widget):
            widget.SetFont(font)
            for child in widget.GetChildren(): set_font_recursive(child)
        set_font_recursive(self.panel)
        self.panel.Layout()
        self.Refresh()
