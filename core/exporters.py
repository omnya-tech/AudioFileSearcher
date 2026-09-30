"""
حفظ نتائج التفريغ بصيغ مختلفة، مستقل عن الواجهة.
المقطع: (نص التوقيت للعرض، النص، البداية، النهاية، قائمة الكلمات بتوقيتاتها)
"""
import json
import re

from core.time_utils import format_range, format_srt_time

FORMATS = ["srt", "txt", "vtt", "json", "docx"]
# علامة BOM في أول ملفات النص والترجمة: بدونها تعرض بعض المشغلات والتلفزيونات القديمة العربي كرموز غريبة
TEXT_ENCODING = "utf-8-sig"

_RTL_CHARS = re.compile("[֐-ࣿיִ-﷿ﹰ-﻿]")


def is_rtl_text(text):
    """النص عربي (أو عبري/فارسي/أردو): اتجاهه من اليمين لليسار، بغض النظر عن لغة الواجهة"""
    return bool(_RTL_CHARS.search(text or ""))


def write_txt(path, segments):
    with open(path, "w", encoding=TEXT_ENCODING) as f:
        f.write("\n\n".join(s[1].strip() for s in segments))


def write_srt(path, segments):
    with open(path, "w", encoding=TEXT_ENCODING) as f:
        for i, s in enumerate(segments, 1):
            f.write(f"{i}\n{format_srt_time(s[2])} --> {format_srt_time(s[3])}\n{s[1]}\n\n")


def write_vtt(path, segments):
    with open(path, "w", encoding=TEXT_ENCODING) as f:
        f.write("WEBVTT\n\n")
        for i, s in enumerate(segments, 1):
            f.write(f"{i}\n{format_srt_time(s[2], '.')} --> {format_srt_time(s[3], '.')}\n{s[1]}\n\n")


def write_json(path, segments):
    data = []
    for s in segments:
        item = {"start": s[2], "end": s[3], "text": s[1]}
        if len(s) > 4 and s[4]:
            item["words"] = s[4]
        data.append(item)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


def write_docx(path, segments, title=""):
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement

    doc = docx.Document()
    if title:
        doc.add_heading(title, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for s in segments:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        rtl = is_rtl_text(s[1])
        if rtl:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            # اتجاه الفقرة من اليمين لليسار حتى يظهر النص العربي مرتباً في الوورد
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        run_time = p.add_run(f"[{format_range(s[2], s[3])}]\n")
        run_time.bold = True
        run_time.font.color.rgb = RGBColor(100, 100, 100)
        run_text = p.add_run(s[1])
        run_text.font.rtl = rtl
    doc.save(path)


def export(path, fmt, segments, title=""):
    """الحفظ بالصيغة المطلوبة (SRT لو الصيغة غير معروفة). ImportError لو مكتبة Word غير مثبتة"""
    if fmt == "docx":
        write_docx(path, segments, title)
    else:
        {"txt": write_txt, "vtt": write_vtt, "json": write_json}.get(fmt, write_srt)(path, segments)
